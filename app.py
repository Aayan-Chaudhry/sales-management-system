from flask import Flask, render_template, request, jsonify, send_file, session, redirect, url_for
import sqlite3, json, io, os, sys, subprocess, tempfile, threading
from datetime import datetime, date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import secrets

app = Flask(__name__)
# New random secret every time the app starts.
# This forces the login page to show again after restarting START.bat instead of remembering the last user.
app.secret_key = os.environ.get('ROUDHA_TRACKER_SECRET_KEY') or secrets.token_hex(32)

ADMIN_PIN = os.environ.get('ROUDHA_ADMIN_PIN', '1234')
EMPLOYEE_PIN = os.environ.get('ROUDHA_EMPLOYEE_PIN', '0000')

PUBLIC_ENDPOINTS = {'login_page', 'login', 'static'}

def current_role():
    return session.get('role')

def require_login():
    if not current_role():
        return redirect(url_for('login_page'))
    return None

def require_admin_api():
    if current_role() != 'admin':
        return jsonify(error='Admin access required.'), 403
    return None

@app.before_request
def protect_pages():
    endpoint = request.endpoint or ''
    if endpoint in PUBLIC_ENDPOINTS or endpoint.startswith('static'):
        return None
    if not current_role():
        if request.path.startswith('/api/'):
            return jsonify(error='Login required.'), 401
        return redirect(url_for('login_page'))
    return None

@app.route('/login')
def login_page():
    # Clear the previous Admin/Employee session whenever the login page opens.
    # This prevents the app from staying logged in after reopening the VBS/app launcher.
    session.clear()
    return '''<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>Roudha Tracker Login</title>
<style>
body{background:#F7F6F2;color:#1A1916;font-family:Georgia,serif;min-height:100vh;display:flex;align-items:center;justify-content:center;margin:0;}
.card{background:#fff;border:1px solid #E2DFD6;border-radius:12px;padding:2rem;width:min(380px,92vw);}
h1{font-size:24px;margin:0 0 1rem;}
label{display:block;font-family:'Courier New',monospace;font-size:12px;color:#7A7870;text-transform:uppercase;margin:.75rem 0 .25rem;}
select,input{width:100%;padding:10px;border:1px solid #E2DFD6;border-radius:7px;font-family:Georgia,serif;font-size:16px;}
button{margin-top:1rem;width:100%;padding:11px;border:0;border-radius:7px;background:#1A1916;color:#F7F6F2;font-family:Georgia,serif;font-size:17px;cursor:pointer;}
.err{color:#A32D2D;margin-top:.75rem;font-family:'Courier New',monospace;font-size:13px;}
.muted{color:#7A7870;font-size:13px;margin-top:.75rem;line-height:1.35;}
</style>
</head>
<body>
<div class="card">
<h1>Roudha Tracker</h1>
<form method="post" action="/api/login">
<label>Access type</label>
<select name="role">
<option value="employee">Employee</option>
<option value="admin">Admin</option>
</select>
<label>PIN</label>
<input name="pin" type="password" autofocus>
<button type="submit">Enter</button>
</form>
<div class="muted">Employee access is limited to New Sale and Labels.</div>
</div>
</body>
</html>'''

@app.route('/api/login', methods=['POST'])
def login():
    data = request.get_json(silent=True) or {}
    role = request.form.get('role') or data.get('role')
    pin = request.form.get('pin') or data.get('pin')
    if role == 'admin' and pin == ADMIN_PIN:
        session['role'] = 'admin'
        return redirect(url_for('index'))
    if role == 'employee' and pin == EMPLOYEE_PIN:
        session['role'] = 'employee'
        return redirect(url_for('index'))
    return '''<!DOCTYPE html><html><body style="font-family:Georgia,serif;background:#F7F6F2;padding:2rem;">
<h2>Wrong PIN</h2><p><a href="/login">Try again</a></p></body></html>''', 401

@app.route('/api/logout', methods=['POST'])
def logout():
    session.clear()
    return jsonify(ok=True)

@app.route('/api/me')
def me():
    return jsonify(role=current_role())

BASE_DIR = os.path.dirname(__file__)
DB = os.path.join(BASE_DIR, 'sales.db')
PTOUCH_CONFIG_FILE = os.path.join(BASE_DIR, 'ptouch_templates.json')
PTOUCH_SCRIPT = os.path.join(BASE_DIR, 'ptouch', 'print_label.py')
PRINT_LOCK = threading.Lock()

# ── INGREDIENTS LIST ──────────────────────────────────────────────────────────
# TO CHANGE: edit this list, save the file, restart the app (re-run START.bat)
INGREDIENTS = [
    "Alth (Althair - Tier 2)",
    "Lay (Layton - Tier 2)",
    "Girl (Good Girl - Tier 3)",
    "Ahoj (Ahojas - Tier 1)",
    "Seven (Oud Seven - Tier 1)",
    "Gypsy (Gypsy Water - Tier 3)",
    "Side (Side Effect - Tier 2)",
    "Love (Psychedelic Love - Tier 2)",
    "Great (Oud for Greatness - Tier 2)",
    "Wanted (Most Wanted - Tier 3)",
    "SWY (Stronger w/ You - Tier 3)",
    "Ghost (Mojave Ghost - Tier 3)",
    "EmOud (Emeraude Oud - Tier 3)",
    "No.77 (Classic No.77 - Tier 3)",
    "Flower (Flowerbomb - Tier 3)",
    "Spice (Spicebomb - Tier 3)",
    "Maliki (Musk Maliki - Tier 3)",
    "Phil (Philosykos - Tier 2)",
    "Son (Do Son - Tier 2)",
    "Icon (Icon Elite - Tier 3)",
    "Nights (New York Nights - Tier 2)",
    "Angel (Angel's Share - Tier 2)",
    "Gio (Gio Profumo - Tier 3)",
    "Mania (Mania Men - Tier 3)",
    "Terre (Terre D'Hermes - Tier 3)",
    "Doxe (Paradoxe - Tier 3)",
    "H24 (Hermes H24 - Tier 3)",
    "Roudha (Roudha Oud - Tier 2)",
    "Blue (Blue Oud - Tier 1)",
    "Qurashi (Qurashi Blend - Tier 2)",
    "OudInt (Oud Intense - Tier 2)",
    "Mazing (Oudmazing - Tier 2)",
    "13 (Another 13 - Tier 1)",
    "49 (Ylang 49 - Tier 2)",
    "29 (Noir 29 - Tier 2)",
    "33 (Santal 33- Tier 2)",
    "Rouge (Rouge 540 - Tier 2)",
    "Soir (Grand Soir - Tier 3)",
    "Mood (Oud Satin Mood - Tier 3)",
    "Imm (L'Immensite - Tier 2)",
    "Imag (Imagination - Tier 2)",
    "Chill (Pacific Chill - Tier 2)",
    "Berg (Oud & Bergamot- Tier 3)",
    "Vet (Vet & Gold V. - Tier 3)",
    "Sage (Sage & Sea Salt - Tier 3)",
    "Peony (Peony & Blush - Tier 3)",
    "Hac (Hacivat - Tier 2)",
    "AvCol (Aventus Cologne - Tier 2)",
    "Av&Ph (Aventus & Ph - Tier 2)",
    "Ani (Ani - Tier 2)",
    "Ever (Forever & Ever - Tier 3)",
    "Chance (Chance - Tier 3)",
    "E.Plat (Egoiste Plat. - Tier 3)",
    "Bleu (Bleu de C. - Tier 2)",
    "Sport (Allure H. Sport- Tier 3)",
    "Dylan (Dylan Blue - Tier 3)",
    "Pour (Pour Homme - Tier 3)",
    "Nomade (Ombre Nomade - Tier 2)",
    "Sable (Sables Roses - Tier 2)",
    "Carbon (Carbon - Tier 2)",
    "Scent (Boss The Scent - Tier 3)",
    "Bottle (Boss Bottled - Tier 3)",
    "D.Int (Dior Homme Intense - Tier 2)",
    "S.Elix (Sauvage Elixir - Tier 2)",
    "S.Parf (Sauvage Parfum - Tier 2)",
    "Sauv (Sauvage - Tier 2)",
    "Halt (Haltane - Tier 1)",
    "Oajan (Oajan - Tier 2)",
    "Naxos (1861 Naxos - Tier 2)",
    "Del (Delina - Tier 2)",
    "Tor21 (Torino 21- Tier 2)",
    "Half (Halfeti - Tier 2)",
    "Y.Elix (YSL Y Elixir - Tier 3)",
    "Libre (Libre - Tier 3)",
    "Y (YSL Y - Tier 3)",
    "Opium (Black Opium - Tier 3)",
    "Myslf (Myslf - Tier 3)",
    "Nuit (Nuit De L'homme - Tier 3)",
    "No.5 (Chanel No.5 - Tier 3)",
    "Fab (F.Fabulous - Tier 3)",
    "Cherry (Lost Cherry - Tier 3)",
    "Wood (Oud Wood - Tier 2)",
    "D.Noir (Noir De Noir - Tier 2)",
    "Leather (Ombre Leather - Tier 2)",
    "Tob.V (tobacco vanille - Tier 3)",
    "Orchid (Black Orchid - Tier 3)",
    "NoirEx (Noir Extreme - Tier 2)",
    "Way (My Way - Tier 3)",
    "One (The One - Tier 3)",
    "Si (Si - Tier 3)",
    "1978 (Polo 1978 - Tier 3)",
    "Marsh (Marshmallow 81 - Tier 2)",
    "Utopia (Utopia V. Coco - Tier 3)",
    "Van.28 (Vanilla 28 - Tier 3)",
    "Chloe (Chloe - Tier 3)",
    "Madem (Coco Mademoiselle - Tier 3)",
    "Daisy (Daisy So Fresh - Tier 3)",
    "Bloom (G.Bloom - Tier 3)",
    "Belle (Vie Est Belle - Tier 3)",
    "Crush (Instant Crush - Tier 3)",
    "Gard. (Gardenia - Tier 3)",
    "Afgano (Black Afgano - Tier 2)",
    "Her (B.Her - Tier 3)",
    "R.musk (Roses & musk - Tier 3)",
    "W.musk (White musk - Tier 2)",
    "Candy (Sweet Like Candy - Tier 3)",
    "Ch.62 (Cheirosa 62 - Tier 3)",
    "Miss D (Miss D - Tier 3)",
    "Goddess (Goddess - Tier 3)",
    "Male (Le Male Elixir - Tier 3)",
    "Eros (Eros - Tier 3)",
    "Roma (Born in Roma - Tier 3)",
    "Olym (Olympea - Tier 3)",
    "Hyp.p (Hypnotic Poison - Tier 3)",
    "Vip.M (212 Vip men - Tier 3)",
    "vip.B (212 Vip Black- Tier 3)",
    "polo.B (Polo Black - Tier 3)",
    "Coral (Coral Fantasy - Tier 3)",
    "Baby (Baby Powder - Tier 2)",
    "Imp.V (Imperial Valley- Tier 2)",
    "Invic (Invictus - Tier 3)",
    "Million (One Million - Tier 3)",
    "Custom",
   # always keep Custom last
]

# ── PRODUCT TYPES ─────────────────────────────────────────────────────────────
PRODUCT_TYPES = [
    "Body Oil",
    "Beard Oil",
    "Oil 500ml",
    "Perfume 20ml",
    "Perfume 50ml",
    "Perfume 100ml",
    "Lotion",
    "Roll On",
    "Diffuser",
    "Custom",
]

# ── PAYMENT METHODS ───────────────────────────────────────────────────────────
PAYMENT_METHODS = [
    "Cash",
    "Card",
    "E-Transfer",
    "Custom",
]

# ── DB SETUP ──────────────────────────────────────────────────────────────────
def get_db():
    conn = sqlite3.connect(DB)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA foreign_keys=ON")
    return conn

def init_db():
    with get_db() as conn:
        conn.executescript("""
        CREATE TABLE IF NOT EXISTS sales (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            customer TEXT NOT NULL,
            sale_date TEXT NOT NULL,
            notes TEXT,
            payments TEXT DEFAULT '[]',
            created_at TEXT DEFAULT (datetime('now','localtime'))
        );
        CREATE TABLE IF NOT EXISTS sale_products (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            sale_id INTEGER NOT NULL REFERENCES sales(id) ON DELETE CASCADE,
            product_type TEXT NOT NULL,
            custom_product_name TEXT,
            product_label TEXT,
            price REAL NOT NULL,
            formula TEXT,
            notes TEXT
        );
        CREATE INDEX IF NOT EXISTS idx_sales_date ON sales(sale_date);
        CREATE TABLE IF NOT EXISTS label_queue (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            sale_id INTEGER NOT NULL REFERENCES sales(id) ON DELETE CASCADE,
            sale_product_id INTEGER,
            customer_name TEXT NOT NULL,
            product_type TEXT NOT NULL,
            product_name TEXT NOT NULL,
            name_on_product TEXT,
            ingredients_text TEXT,
            template_file TEXT,
            status TEXT DEFAULT 'not_printed',
            error TEXT,
            created_at TEXT DEFAULT (datetime('now','localtime')),
            printed_at TEXT
        );
        CREATE INDEX IF NOT EXISTS idx_label_queue_status ON label_queue(status);
        CREATE INDEX IF NOT EXISTS idx_label_queue_sale ON label_queue(sale_id);
        """)
        # Existing users may already have a sales table without the payments column.
        cols = [r[1] for r in conn.execute("PRAGMA table_info(sales)").fetchall()]
        if 'payments' not in cols:
            conn.execute("ALTER TABLE sales ADD COLUMN payments TEXT DEFAULT '[]'")

        product_cols = [r[1] for r in conn.execute("PRAGMA table_info(sale_products)").fetchall()]
        if 'product_label' not in product_cols:
            conn.execute("ALTER TABLE sale_products ADD COLUMN product_label TEXT DEFAULT ''")
        if 'essence_no' not in product_cols:
            conn.execute("ALTER TABLE sale_products ADD COLUMN essence_no TEXT DEFAULT ''")
        if 'tier' not in product_cols:
            conn.execute("ALTER TABLE sale_products ADD COLUMN tier TEXT DEFAULT ''")

        label_cols = [r[1] for r in conn.execute("PRAGMA table_info(label_queue)").fetchall()]
        if 'essence_no' not in label_cols:
            conn.execute("ALTER TABLE label_queue ADD COLUMN essence_no TEXT DEFAULT ''")
        if 'tier' not in label_cols:
            conn.execute("ALTER TABLE label_queue ADD COLUMN tier TEXT DEFAULT ''")

init_db()


# ── FORMULA VALIDATION ─────────────────────────────────────────────────────────
def validate_product_formula(product, product_index=1):
    """Validate one product before saving.

    Frontend checks help the user, but backend validation is the real guard.
    Every saved product must have ingredients that total exactly 100%.
    """
    product_type = (product.get('product_type') or '').strip()
    custom_product_name = (product.get('custom_product_name') or '').strip()
    product_label = (product.get('product_label') or '').strip()
    essence_no = (product.get('essence_no') or '').strip()
    tier = (product.get('tier') or '').strip()
    label = custom_product_name if product_type == 'Custom' and custom_product_name else (product_type or f'Product {product_index}')

    if not product_type:
        raise ValueError(f'Product {product_index} is missing a product type.')
    if not product_label:
        raise ValueError(f'"{label}" is missing Name on Product.')
    if product_type == 'Custom' and not custom_product_name:
        raise ValueError('Custom product is missing a product name.')
    if product_type == 'Oil 500ml' and not essence_no:
        raise ValueError('Oil 500ml is missing Essence NO.')
    if product_type == 'Oil 500ml' and not tier:
        raise ValueError('Oil 500ml is missing Tier.')

    if product_type == 'Oil 500ml':
        # Oil 500ml uses the Price box position for Tier instead.
        # Keep sale revenue/payment math safe by storing price as 0.
        price = 0.0
    else:
        try:
            price = float(product.get('price'))
        except (TypeError, ValueError):
            raise ValueError(f'"{label}" is missing a valid price.')

    formula = product.get('formula', [])
    if not isinstance(formula, list) or not formula:
        raise ValueError(f'"{label}" must have at least one ingredient.')

    cleaned_formula = []
    total_pct = 0.0

    for row_index, item in enumerate(formula, start=1):
        if not isinstance(item, dict):
            raise ValueError(f'"{label}" has an invalid ingredient row.')

        name = (item.get('name') or '').strip()
        try:
            pct = float(item.get('pct'))
        except (TypeError, ValueError):
            raise ValueError(f'"{label}" ingredient row {row_index} is missing a valid percentage.')

        if not name:
            raise ValueError(f'"{label}" ingredient row {row_index} is missing an ingredient name.')
        if pct <= 0:
            raise ValueError(f'"{label}" ingredient "{name}" must be above 0%.')

        pct = round(pct, 1)
        total_pct += pct
        cleaned_formula.append({'name': name, 'pct': pct})

    total_pct = round(total_pct, 1)
    if total_pct != 100.0:
        raise ValueError(f'"{label}" ingredients add to {total_pct}% — they must total exactly 100%.')

    return {
        'product_type': product_type,
        'custom_product_name': custom_product_name,
        'product_label': product_label,
        'essence_no': essence_no,
        'tier': tier,
        'price': price,
        'formula': cleaned_formula,
        'notes': product.get('notes', '')
    }

def validate_products(products):
    if not isinstance(products, list) or not products:
        raise ValueError('At least one product required.')
    return [validate_product_formula(p, i) for i, p in enumerate(products, start=1)]


def ingredient_print_name(name, single_ingredient=False):
    """Return the ingredient short form exactly how it should print on P-touch labels.

    Rule:
    - Always use the short form before parentheses for ingredients.
      Example: Alth (Althair - Tier 3) -> Alth
    - This applies whether there is 1 ingredient or multiple ingredients.
    """
    text = (name or '').strip()

    if '(' in text and ')' in text:
        before = text.split('(', 1)[0].strip()
        return before or text

    return text


def format_formula_for_label(formula):
    """Turn the saved formula list into clean label text for P-touch.

    Label rule:
    - Ingredient names always print using the short form before parentheses.
      Example: Lay (Layton - Tier 2) -> Lay
    - Even split mixes print without percentages.
      Example: Lay + Rouge -> Lay+Rouge
    - Uneven custom mixes print with percentages.
      Example: Lay 30% + Rouge 70% -> Lay(30)+Rouge(70)
    """
    if not formula:
        return ''

    cleaned = []
    for item in formula:
        name = (item.get('name') or '').strip()
        if not name:
            continue
        try:
            pct = float(item.get('pct') or 0)
            pct_text = str(int(pct)) if pct.is_integer() else f"{pct:.1f}".rstrip('0').rstrip('.')
        except Exception:
            pct = None
            pct_text = str(item.get('pct') or '').strip()
        cleaned.append((name, pct, pct_text))

    if not cleaned:
        return ''

    def is_even_split(rows):
        """True when percentages match the app's auto-even split.

        Uses tenths of a percent so 3-way mixes like 33.4/33.3/33.3 count as even.
        """
        if not rows:
            return False
        actual = []
        for _name, pct, _pct_text in rows:
            if pct is None:
                return False
            actual.append(int(round(float(pct) * 10)))

        n = len(actual)
        total_tenths = 1000
        base = total_tenths // n
        remainder = total_tenths % n
        expected = [base + (1 if i < remainder else 0) for i in range(n)]
        return sorted(actual) == sorted(expected)

    short_names = [ingredient_print_name(name, single_ingredient=(len(cleaned) == 1)) for name, _pct, _pct_text in cleaned]

    # One 100% ingredient or any even split mix prints names only, no percentages.
    if is_even_split(cleaned):
        return '+'.join(short_names)

    # Uneven custom percentages print with percentages.
    parts = []
    for (name, pct, pct_text), short_name in zip(cleaned, short_names):
        parts.append(f"{short_name}({pct_text})" if pct_text else short_name)
    return '+'.join(parts)


def product_display_name(product):
    """Return the product name that should appear on reports/labels."""
    product_type = (product.get('product_type') or '').strip()
    custom_name = (product.get('custom_product_name') or '').strip()
    if product_type == 'Custom' and custom_name:
        return custom_name
    return product_type or custom_name or 'Product'


def load_ptouch_config():
    default = {
        'printing_enabled': False,
        'simulate_printing': True,
        'printer_name': '',
        'templates': {}
    }
    if not os.path.exists(PTOUCH_CONFIG_FILE):
        return default
    try:
        with open(PTOUCH_CONFIG_FILE, 'r', encoding='utf-8') as f:
            data = json.load(f)
        default.update(data or {})
        default['templates'] = data.get('templates', {}) if isinstance(data, dict) else {}
    except Exception:
        pass
    return default


def resolve_ptouch_template(product):
    cfg = load_ptouch_config()
    templates = cfg.get('templates', {}) or {}
    product_type = (product.get('product_type') or '').strip()
    template = templates.get(product_type)
    if product_type == 'Custom':
        template = templates.get('Custom') or template
    if not template:
        return ''
    if not os.path.isabs(template):
        template = os.path.join(BASE_DIR, template)
    return os.path.normpath(template)


def delete_existing_label_jobs_for_sale_product(conn, sale_product_id):
    """Delete stale unprinted jobs before creating a fresh product label job."""
    if not sale_product_id:
        return
    conn.execute(
        "DELETE FROM label_queue WHERE sale_product_id=? AND status IN ('not_printed','error','simulated','printing')",
        (sale_product_id,)
    )



def delete_existing_label_jobs_for_sale(conn, sale_id):
    """Delete stale non-final label jobs for a sale before rebuilding fresh jobs.

    Printed labels stay as history. Not-printed/error/simulated/printing rows are
    old work-in-progress queue rows and must not be reused after a sale edit.
    """
    if not sale_id:
        return
    conn.execute(
        "DELETE FROM label_queue WHERE sale_id=? AND status IN ('not_printed','error','simulated','printing')",
        (sale_id,)
    )


def product_row_to_label_product(product_row):
    """Build the print product payload from the current saved sale_products row."""
    return {
        'product_type': product_row['product_type'],
        'custom_product_name': product_row['custom_product_name'],
        'product_label': product_row['product_label'] if 'product_label' in product_row.keys() else '',
        'essence_no': product_row['essence_no'] if 'essence_no' in product_row.keys() else '',
        'tier': product_row['tier'] if 'tier' in product_row.keys() else '',
        'price': product_row['price'],
        'formula': json.loads(product_row['formula']) if product_row['formula'] else [],
        'notes': product_row['notes'],
    }


def create_label_queue_item(conn, sale_id, sale_product_id, customer, product):
    # Custom products should be saved in the tracker only.
    # Do not create a P-touch label job for them.
    if (product.get('product_type') or '').strip() == 'Custom':
        return False

    delete_existing_label_jobs_for_sale_product(conn, sale_product_id)

    product_name = product_display_name(product)
    name_on_product = (product.get('product_label') or '').strip() or customer
    ingredients_text = format_formula_for_label(product.get('formula') or [])
    essence_no = (product.get('essence_no') or '').strip()
    tier = (product.get('tier') or '').strip()
    template_file = resolve_ptouch_template(product)
    cur = conn.execute("""
        INSERT INTO label_queue
        (sale_id, sale_product_id, customer_name, product_type, product_name, name_on_product, ingredients_text, essence_no, tier, template_file)
        VALUES (?,?,?,?,?,?,?,?,?,?)
    """, (
        sale_id, sale_product_id, customer, product.get('product_type',''), product_name,
        name_on_product, ingredients_text, essence_no, tier, template_file
    ))
    return cur.lastrowid


def label_row_to_job(row, chain_printing=False):
    cfg = load_ptouch_config()
    fields = {
        'NAME_ON_PRODUCT': row['name_on_product'] or '',
        'INGREDIENTS': row['ingredients_text'] or '',
    }

    # ESSENCE_NO is only for the Oil 500ml template.
    # Do not send it to 20ml/50ml/100ml/etc., because those P-touch templates
    # do not have an ESSENCE_NO object and b-PAC will error if an object is missing.
    if (row['product_type'] or '').strip() == 'Oil 500ml':
        fields['ESSENCE_NO'] = row['essence_no'] or ''
        fields['TIER'] = row['tier'] or ''

    return {
        'label_id': row['id'],
        'template': row['template_file'],
        'printer_name': cfg.get('printer_name', ''),
        'simulate_printing': bool(cfg.get('simulate_printing', True)) or not bool(cfg.get('printing_enabled', False)),
        'chain_printing': bool(chain_printing),
        'fields': fields
    }


def run_ptouch_print_job(job):
    """Run the Windows b-PAC print script in a subprocess.

    On non-Windows/dev machines, or when simulate_printing is true, the script writes a preview JSON
    instead of trying to access Brother b-PAC.
    """
    os.makedirs(os.path.join(BASE_DIR, 'ptouch', 'jobs'), exist_ok=True)
    fd, job_path = tempfile.mkstemp(prefix='ptouch_job_', suffix='.json', dir=os.path.join(BASE_DIR, 'ptouch', 'jobs'))
    with os.fdopen(fd, 'w', encoding='utf-8') as f:
        json.dump(job, f, indent=2)
    cmd = [sys.executable, PTOUCH_SCRIPT, '--job', job_path]
    proc = subprocess.run(cmd, cwd=BASE_DIR, capture_output=True, text=True, timeout=60)
    if proc.returncode != 0:
        raise RuntimeError((proc.stderr or proc.stdout or 'P-touch print script failed.').strip())
    return (proc.stdout or '').strip()


def print_label_queue_row(conn, label_id, chain_printing=False, allow_reprint=False):
    """Print one queued label safely.

    Safety rules:
    - A printed label cannot be printed again from the queue unless code explicitly allows it.
    - A label already marked as printing cannot be started again.
    - The row is marked as printing before Brother b-PAC is called, then committed,
      so double-clicks or repeated requests cannot trigger the same job twice.
    """
    with PRINT_LOCK:
        row = conn.execute("SELECT * FROM label_queue WHERE id=?", (label_id,)).fetchone()
        if not row:
            raise ValueError('Label not found.')

        current_status = (row['status'] or 'not_printed').strip()
        if current_status == 'printed' and not allow_reprint:
            raise ValueError('This label is already printed. Create a fresh label job if you need a reprint.')
        if current_status == 'printing':
            raise ValueError('This label is already printing. Wait for it to finish.')

        conn.execute(
            "UPDATE label_queue SET status='printing', error='' WHERE id=?",
            (label_id,)
        )
        conn.commit()

    job = label_row_to_job(row, chain_printing=chain_printing)

    try:
        output = run_ptouch_print_job(job)
        status = 'printed' if not job.get('simulate_printing') else 'simulated'
        conn.execute(
            "UPDATE label_queue SET status=?, error='', printed_at=datetime('now','localtime') WHERE id=?",
            (status, label_id)
        )
        return output, status
    except Exception:
        conn.execute(
            "UPDATE label_queue SET status='error', printed_at=NULL WHERE id=?",
            (label_id,)
        )
        raise


def validate_payments(payments, expected_total):
    """Validate sale-level payments. Supports split payments like Cash + Card.

    Payment total must match product total so revenue stays trustworthy.
    """
    if not isinstance(payments, list) or not payments:
        raise ValueError('Add at least one payment method.')

    cleaned = []
    total_paid = 0.0
    for idx, pay in enumerate(payments, start=1):
        if not isinstance(pay, dict):
            raise ValueError(f'Payment row {idx} is invalid.')

        method = (pay.get('method') or '').strip()
        custom_method = (pay.get('custom_method') or '').strip()
        label = custom_method if method == 'Custom' and custom_method else method

        if not method:
            raise ValueError(f'Payment row {idx} is missing a payment method.')
        if method == 'Custom' and not custom_method:
            raise ValueError('Custom payment method is missing a name.')

        try:
            amount = float(pay.get('amount'))
        except (TypeError, ValueError):
            raise ValueError(f'Payment row {idx} is missing a valid amount.')

        if amount < 0:
            raise ValueError(f'Payment amount for "{label}" cannot be negative.')

        # $0 payments are allowed, for cases like free items, samples, exchanges,
        # discounts, or fully-comped sales. Negative payments are still blocked.
        amount = round(amount, 2)

        total_paid += amount
        cleaned.append({'method': method, 'custom_method': custom_method, 'amount': amount})

    total_paid = round(total_paid, 2)
    expected_total = round(float(expected_total or 0), 2)
    if total_paid != expected_total:
        raise ValueError(f'Payment total is ${total_paid:.2f}, but sale total is ${expected_total:.2f}. Payment amounts must match the sale total.')

    return cleaned

def parse_payments(value):
    """Return payments as a clean list, even if the DB has old/blank/string data."""
    if not value:
        return []

    if isinstance(value, list):
        data = value
    elif isinstance(value, str):
        try:
            data = json.loads(value)
        except Exception:
            # Old/non-JSON text should not crash exports. Show it as a note-like payment.
            return [{'method': value, 'custom_method': '', 'amount': 0}] if value.strip() else []
    else:
        return []

    if isinstance(data, dict):
        data = [data]
    if not isinstance(data, list):
        return []

    cleaned = []
    for p in data:
        if isinstance(p, dict):
            cleaned.append(p)
        elif isinstance(p, str) and p.strip():
            cleaned.append({'method': p.strip(), 'custom_method': '', 'amount': 0})
    return cleaned

def payment_summary(payments):
    """Human-friendly payment text for app display and Excel exports.

    Single-method sales show just the method, e.g. "Cash".
    Split payments show amounts, e.g. "Cash: $20.00 + Card: $5.00".
    """
    payments = parse_payments(payments)
    if not payments:
        return ''

    def label_and_amount(p):
        method = p.get('custom_method') if p.get('method') == 'Custom' and p.get('custom_method') else p.get('method', '')
        try:
            amount = float(p.get('amount') or 0)
        except (TypeError, ValueError):
            amount = 0
        return str(method), amount

    if len(payments) == 1:
        method, _amount = label_and_amount(payments[0])
        return method

    parts = []
    for p in payments:
        method, amount = label_and_amount(p)
        if not method:
            continue
        parts.append(f"{method}: ${amount:.2f}" if amount else method)
    return ' + '.join(parts)


# ── ROUTES ────────────────────────────────────────────────────────────────────
@app.route('/')
def index():
    return render_template('index.html',
        ingredients=INGREDIENTS,
        product_types=PRODUCT_TYPES,
        payment_methods=PAYMENT_METHODS,
        role=current_role())

@app.route('/api/config')
def config():
    return jsonify(ingredients=INGREDIENTS, product_types=PRODUCT_TYPES, payment_methods=PAYMENT_METHODS)

@app.route('/api/sales', methods=['POST'])
def create_sale():
    data = request.json or {}
    customer = (data.get('customer') or '').strip()
    if not customer or not data.get('products'):
        return jsonify(error='Customer name and at least one product required.'), 400

    try:
        products = validate_products(data.get('products'))
        sale_total = round(sum(float(p['price']) for p in products), 2)
        payments = validate_payments(data.get('payments'), sale_total)
    except ValueError as e:
        return jsonify(error=str(e)), 400

    sale_date = data.get('sale_date') or date.today().isoformat()
    with get_db() as conn:
        cur = conn.execute(
            "INSERT INTO sales (customer, sale_date, notes, payments) VALUES (?,?,?,?)",
            (customer, sale_date, data.get('notes',''), json.dumps(payments))
        )
        sale_id = cur.lastrowid
        label_count = 0
        for p in products:
            prod_cur = conn.execute(
                "INSERT INTO sale_products (sale_id, product_type, custom_product_name, product_label, essence_no, tier, price, formula, notes) VALUES (?,?,?,?,?,?,?,?,?)",
                (sale_id, p['product_type'], p.get('custom_product_name',''), p.get('product_label',''), p.get('essence_no',''), p.get('tier',''), p['price'],
                 json.dumps(p['formula']), p.get('notes',''))
            )
            if create_label_queue_item(conn, sale_id, prod_cur.lastrowid, customer, p):
                label_count += 1
    return jsonify(id=sale_id, ok=True, labels_added=label_count)

@app.route('/api/stats')
def get_stats():
    guard = require_admin_api()
    if guard: return guard
    """Lightweight dashboard stats — pure SQL aggregates, never loads full records."""
    today_str = date.today().isoformat()
    with get_db() as conn:
        today_row = conn.execute(
            "SELECT COUNT(sp.id) as products, COALESCE(SUM(sp.price),0) as revenue "
            "FROM sales s JOIN sale_products sp ON sp.sale_id=s.id "
            "WHERE s.sale_date=?", (today_str,)).fetchone()
        all_row = conn.execute(
            "SELECT COUNT(sp.id) as products, COALESCE(SUM(sp.price),0) as revenue "
            "FROM sale_products sp").fetchone()
        return jsonify(
            today_products=today_row['products'],
            today_revenue=round(today_row['revenue'], 2),
            all_products=all_row['products'],
            all_revenue=round(all_row['revenue'], 2),
        )

@app.route('/api/reports')
def get_reports():
    guard = require_admin_api()
    if guard: return guard
    """All report aggregations done in SQL — scales to any number of records."""
    with get_db() as conn:
        # Top products by revenue
        top_products = conn.execute("""
            SELECT CASE WHEN sp.product_type='Custom' AND COALESCE(sp.custom_product_name,'') != '' THEN 'Custom: "' || sp.custom_product_name || '"' ELSE sp.product_type END as name,
                   COUNT(*) as cnt, ROUND(SUM(sp.price),2) as rev
            FROM sale_products sp
            GROUP BY name ORDER BY rev DESC LIMIT 20
        """).fetchall()
        # Top customers by revenue
        top_customers = conn.execute("""
            SELECT s.customer, COUNT(DISTINCT s.id) as cnt,
                   ROUND(SUM(sp.price),2) as rev
            FROM sales s JOIN sale_products sp ON sp.sale_id=s.id
            GROUP BY s.customer ORDER BY rev DESC LIMIT 20
        """).fetchall()
        # Ingredient revenue — weighted by pct share stored in formula JSON
        # We do this in Python since formula is JSON; but we only load formula data, not full records
        ing_rows = conn.execute(
            "SELECT price, formula FROM sale_products WHERE formula IS NOT NULL AND formula != '[]'"
        ).fetchall()
        by_ing = {}
        for row in ing_rows:
            try:
                formula = json.loads(row['formula'])
                total_pct = sum(f['pct'] for f in formula)
                for f in formula:
                    name = f['name']
                    if not name: continue
                    if name not in by_ing:
                        by_ing[name] = {'count': 0, 'rev': 0.0}
                    by_ing[name]['count'] += 1
                    by_ing[name]['rev'] += row['price'] * (f['pct'] / total_pct) if total_pct > 0 else 0
            except Exception:
                pass
        top_ing = sorted(by_ing.items(), key=lambda x: -x[1]['rev'])[:20]
        return jsonify(
            top_products=[{'name': r['name'], 'count': r['cnt'], 'rev': r['rev']} for r in top_products],
            top_customers=[{'name': r['customer'], 'count': r['cnt'], 'rev': r['rev']} for r in top_customers],
            top_ingredients=[{'name': n, 'count': d['count'], 'rev': round(d['rev'], 2)} for n, d in top_ing],
        )

def _fetch_sales_with_products(conn, rows):
    """Helper: attach products to a list of sale rows."""
    if not rows:
        return []
    sale_ids = [s['id'] for s in rows]
    placeholders = ','.join('?' * len(sale_ids))
    prods = conn.execute(
        f"SELECT * FROM sale_products WHERE sale_id IN ({placeholders}) ORDER BY sale_id, id",
        sale_ids).fetchall()
    prod_map = {}
    for p in prods:
        prod_map.setdefault(p['sale_id'], []).append({
            'id': p['id'], 'product_type': p['product_type'],
            'custom_product_name': p['custom_product_name'],
            'product_label': p['product_label'] if 'product_label' in p.keys() else '',
            'essence_no': p['essence_no'] if 'essence_no' in p.keys() else '',
            'tier': p['tier'] if 'tier' in p.keys() else '',
            'price': p['price'],
            'formula': json.loads(p['formula']) if p['formula'] else [],
            'notes': p['notes']
        })
    return [{
        'id': s['id'], 'customer': s['customer'],
        'sale_date': s['sale_date'], 'notes': s['notes'],
        'payments': parse_payments(s['payments']),
        'created_at': s['created_at'],
        'products': prod_map.get(s['id'], [])
    } for s in rows]

PAGE_SIZE = 50

@app.route('/api/sales', methods=['GET'])
def get_sales():
    date_filter = request.args.get('date')
    if current_role() != 'admin':
        if date_filter != date.today().isoformat():
            return jsonify(error='Admin access required.'), 403
    month_filter = request.args.get('month')
    year_filter = request.args.get('year')
    search = request.args.get('search', '').strip()
    page = max(1, int(request.args.get('page', 1)))
    offset = (page - 1) * PAGE_SIZE

    with get_db() as conn:
        if date_filter:
            # Single day — always small, no pagination needed
            rows = conn.execute(
                "SELECT * FROM sales WHERE sale_date=? ORDER BY created_at DESC",
                (date_filter,)).fetchall()
            return jsonify(_fetch_sales_with_products(conn, rows))

        if month_filter:
            rows = conn.execute(
                "SELECT * FROM sales WHERE sale_date LIKE ? ORDER BY sale_date DESC, created_at DESC",
                (month_filter + '%',)).fetchall()
            return jsonify(_fetch_sales_with_products(conn, rows))

        if year_filter:
            rows = conn.execute(
                "SELECT * FROM sales WHERE sale_date LIKE ? ORDER BY sale_date DESC, created_at DESC",
                (year_filter + '%',)).fetchall()
            return jsonify(_fetch_sales_with_products(conn, rows))

        if search:
            rows = conn.execute(
                "SELECT * FROM sales WHERE customer LIKE ? ORDER BY sale_date DESC, created_at DESC",
                (f'%{search}%',)).fetchall()
            return jsonify(_fetch_sales_with_products(conn, rows))

        # Default (History tab) — paginated
        total = conn.execute("SELECT COUNT(*) FROM sales").fetchone()[0]
        rows = conn.execute(
            "SELECT * FROM sales ORDER BY sale_date DESC, created_at DESC LIMIT ? OFFSET ?",
            (PAGE_SIZE, offset)).fetchall()
        return jsonify({
            'sales': _fetch_sales_with_products(conn, rows),
            'total': total,
            'page': page,
            'page_size': PAGE_SIZE,
            'pages': (total + PAGE_SIZE - 1) // PAGE_SIZE,
        })


@app.route('/api/sales/today', methods=['GET'])
def get_today_sales():
    today_str = date.today().isoformat()
    with get_db() as conn:
        rows = conn.execute(
            "SELECT * FROM sales WHERE sale_date=? ORDER BY created_at DESC",
            (today_str,)
        ).fetchall()
        return jsonify(_fetch_sales_with_products(conn, rows))


@app.route('/api/sales/<int:sale_id>', methods=['GET'])
def get_sale(sale_id):
    if current_role() != 'admin':
        with get_db() as conn:
            sale = conn.execute("SELECT sale_date FROM sales WHERE id=?", (sale_id,)).fetchone()
        if not sale or sale['sale_date'] != date.today().isoformat():
            return jsonify(error='Employees can only view today\'s sales.'), 403
    with get_db() as conn:
        s = conn.execute("SELECT * FROM sales WHERE id=?", (sale_id,)).fetchone()
        if not s: return jsonify(error='Not found'), 404
        prods = conn.execute("SELECT * FROM sale_products WHERE sale_id=?", (sale_id,)).fetchall()
        return jsonify({
            'id': s['id'], 'customer': s['customer'],
            'sale_date': s['sale_date'], 'notes': s['notes'],
            'payments': parse_payments(s['payments']),
            'products': [{
                'id': p['id'], 'product_type': p['product_type'],
                'custom_product_name': p['custom_product_name'],
                'product_label': p['product_label'] if 'product_label' in p.keys() else '',
                'essence_no': p['essence_no'] if 'essence_no' in p.keys() else '',
                'tier': p['tier'] if 'tier' in p.keys() else '',
                'price': p['price'],
                'formula': json.loads(p['formula']) if p['formula'] else [],
                'notes': p['notes']
            } for p in prods]
        })

@app.route('/api/sales/<int:sale_id>', methods=['PUT'])
def update_sale(sale_id):
    data = request.json or {}
    if current_role() != 'admin':
        with get_db() as conn:
            sale = conn.execute("SELECT sale_date FROM sales WHERE id=?", (sale_id,)).fetchone()
        today_str = date.today().isoformat()
        if not sale or sale['sale_date'] != today_str or data.get('sale_date') != today_str:
            return jsonify(error='Employees can only edit today\'s sales.'), 403
    customer = (data.get('customer') or '').strip()
    if not customer:
        return jsonify(error='Customer name required.'), 400

    try:
        products = validate_products(data.get('products'))
        sale_total = round(sum(float(p['price']) for p in products), 2)
        payments = validate_payments(data.get('payments'), sale_total)
    except ValueError as e:
        return jsonify(error=str(e)), 400

    with get_db() as conn:
        conn.execute("UPDATE sales SET customer=?, sale_date=?, notes=?, payments=? WHERE id=?",
            (customer, data['sale_date'], data.get('notes',''), json.dumps(payments), sale_id))

        # Editing a sale should update the tracker only.
        # Do NOT delete/recreate label queue items here, because users may open Edit
        # just to review or correct sale details and should not accidentally send
        # the products back into the label printing queue.
        conn.execute("DELETE FROM sale_products WHERE sale_id=?", (sale_id,))
        for p in products:
            conn.execute(
                "INSERT INTO sale_products (sale_id, product_type, custom_product_name, product_label, essence_no, tier, price, formula, notes) VALUES (?,?,?,?,?,?,?,?,?)",
                (sale_id, p['product_type'], p.get('custom_product_name',''), p.get('product_label',''), p.get('essence_no',''), p.get('tier',''), p['price'],
                 json.dumps(p['formula']), p.get('notes',''))
            )
    return jsonify(ok=True, labels_added=0)

@app.route('/api/sales/<int:sale_id>', methods=['DELETE'])
def delete_sale(sale_id):
    if current_role() != 'admin':
        return jsonify(error='Admin access required.'), 403
    with get_db() as conn:
        conn.execute("DELETE FROM sales WHERE id=?", (sale_id,))
    return jsonify(ok=True)

@app.route('/api/labels', methods=['GET'])
def get_label_queue():
    status = request.args.get('status', '').strip()
    limit = max(1, min(200, int(request.args.get('limit', 100))))
    with get_db() as conn:
        if status:
            rows = conn.execute(
                "SELECT * FROM label_queue WHERE status=? ORDER BY created_at DESC, id DESC LIMIT ?",
                (status, limit)
            ).fetchall()
        else:
            rows = conn.execute(
                "SELECT * FROM label_queue ORDER BY created_at DESC, id DESC LIMIT ?",
                (limit,)
            ).fetchall()
        return jsonify([dict(r) for r in rows])


@app.route('/api/labels/<int:label_id>/print', methods=['POST'])
def print_one_label(label_id):
    with get_db() as conn:
        try:
            output, status = print_label_queue_row(conn, label_id)
            return jsonify(ok=True, status=status, output=output)
        except Exception as e:
            conn.execute("UPDATE label_queue SET status='error', error=? WHERE id=?", (str(e), label_id))
            return jsonify(error=str(e)), 500


@app.route('/api/sales/<int:sale_id>/products/<int:product_id>/label/print', methods=['POST'])
def print_sale_product_label(sale_id, product_id):
    with get_db() as conn:
        sale = conn.execute("SELECT * FROM sales WHERE id=?", (sale_id,)).fetchone()
        if not sale:
            return jsonify(error='Sale not found.'), 404

        product_row = conn.execute(
            "SELECT * FROM sale_products WHERE id=? AND sale_id=?",
            (product_id, sale_id)
        ).fetchone()
        if not product_row:
            return jsonify(error='Product not found for this sale.'), 404

        product = {
            'product_type': product_row['product_type'],
            'custom_product_name': product_row['custom_product_name'],
            'product_label': product_row['product_label'] if 'product_label' in product_row.keys() else '',
            'essence_no': product_row['essence_no'] if 'essence_no' in product_row.keys() else '',
            'tier': product_row['tier'] if 'tier' in product_row.keys() else '',
            'price': product_row['price'],
            'formula': json.loads(product_row['formula']) if product_row['formula'] else [],
            'notes': product_row['notes'],
        }

        if (product.get('product_type') or '').strip() == 'Custom':
            return jsonify(error='Custom products do not have a P-touch label template.'), 400

        label_id = create_label_queue_item(conn, sale_id, product_id, sale['customer'], product)
        if not label_id:
            return jsonify(error='Could not create label job for this product.'), 400

        try:
            output, status = print_label_queue_row(conn, label_id)
            return jsonify(ok=True, label_id=label_id, status=status, output=output)
        except Exception as e:
            conn.execute("UPDATE label_queue SET status='error', error=? WHERE id=?", (str(e), label_id))
            return jsonify(error=str(e), label_id=label_id), 500


@app.route('/api/sales/<int:sale_id>/labels/print', methods=['POST'])
def print_sale_labels(sale_id):
    results = []
    with get_db() as conn:
        sale = conn.execute("SELECT * FROM sales WHERE id=?", (sale_id,)).fetchone()
        if not sale:
            return jsonify(error='Sale not found.'), 404

        # CRITICAL FIX:
        # Do not reuse label_queue rows that were made when the sale was first saved.
        # If the sale was edited after that, those rows can be missing added products
        # or still include removed/old products. Rebuild from current sale_products every time.
        delete_existing_label_jobs_for_sale(conn, sale_id)

        product_rows = conn.execute(
            "SELECT * FROM sale_products WHERE sale_id=? ORDER BY id",
            (sale_id,)
        ).fetchall()

        rows = []
        for product_row in product_rows:
            product = product_row_to_label_product(product_row)
            if (product.get('product_type') or '').strip() == 'Custom':
                continue
            label_id = create_label_queue_item(conn, sale_id, product_row['id'], sale['customer'], product)
            if label_id:
                rows.append({'id': label_id})

        total_labels = len(rows)
        for index, row in enumerate(rows):
            try:
                chain_printing = total_labels > 1 and index < total_labels - 1
                output, status = print_label_queue_row(conn, row['id'], chain_printing=chain_printing)
                results.append({'id': row['id'], 'ok': True, 'status': status, 'output': output, 'chain_printing': chain_printing})
            except Exception as e:
                conn.execute("UPDATE label_queue SET status='error', error=? WHERE id=?", (str(e), row['id']))
                results.append({'id': row['id'], 'ok': False, 'error': str(e)})
    return jsonify(ok=all(r.get('ok') for r in results), results=results)


@app.route('/api/labels/print-unprinted', methods=['POST'])
def print_all_unprinted_labels():
    results = []
    with get_db() as conn:
        rows = conn.execute(
            "SELECT id FROM label_queue WHERE status IN ('not_printed','error') ORDER BY id LIMIT 100"
        ).fetchall()
        total_labels = len(rows)
        for index, row in enumerate(rows):
            try:
                chain_printing = total_labels > 1 and index < total_labels - 1
                output, status = print_label_queue_row(conn, row['id'], chain_printing=chain_printing)
                results.append({'id': row['id'], 'ok': True, 'status': status, 'output': output, 'chain_printing': chain_printing})
            except Exception as e:
                conn.execute("UPDATE label_queue SET status='error', error=? WHERE id=?", (str(e), row['id']))
                results.append({'id': row['id'], 'ok': False, 'error': str(e)})
    return jsonify(ok=all(r.get('ok') for r in results), results=results)


@app.route('/api/labels/<int:label_id>/reset', methods=['POST'])
def reset_label(label_id):
    with get_db() as conn:
        row = conn.execute("SELECT status FROM label_queue WHERE id=?", (label_id,)).fetchone()
        if not row:
            return jsonify(error='Label not found.'), 404
        if row['status'] in ('printed','printing'):
            return jsonify(error='Printed or currently-printing labels cannot be reset.'), 400
        conn.execute("UPDATE label_queue SET status='not_printed', error='', printed_at=NULL WHERE id=?", (label_id,))
    return jsonify(ok=True)


@app.route('/api/labels/<int:label_id>', methods=['DELETE'])
def delete_label(label_id):
    with get_db() as conn:
        row = conn.execute("SELECT status FROM label_queue WHERE id=?", (label_id,)).fetchone()
        if not row:
            return jsonify(error='Label not found.'), 404
        if row['status'] in ('printed','printing'):
            return jsonify(error='Printed or currently-printing labels cannot be deleted.'), 400
        conn.execute("DELETE FROM label_queue WHERE id=?", (label_id,))
    return jsonify(ok=True)


@app.route('/api/days')
def get_days():
    guard = require_admin_api()
    if guard: return guard
    with get_db() as conn:
        rows = conn.execute(
            "SELECT sale_date, COUNT(*) as sales, SUM(sp.price) as revenue "
            "FROM sales s JOIN sale_products sp ON sp.sale_id=s.id "
            "GROUP BY sale_date ORDER BY sale_date DESC").fetchall()
        return jsonify([dict(r) for r in rows])

@app.route('/api/export')
def export_excel():
    guard = require_admin_api()
    if guard: return guard
    """Export sales to Excel.

    - mode=year: one workbook for a selected year, with 12 monthly sheets plus a YEAR TOTAL sheet.
      Each monthly sheet groups rows by sale day, shows a DAY TOTAL after each day, and a MONTH TOTAL at the bottom.
    - mode=day: keeps the old single-day export for quick daily downloads.
    """
    mode = request.args.get('mode', 'year')  # 'year' or 'day'
    day = request.args.get('date', date.today().isoformat())
    year = str(request.args.get('year', date.today().year)).strip()

    if not year.isdigit() or len(year) != 4:
        year = str(date.today().year)

    with get_db() as conn:
        if mode == 'day':
            sales_rows = conn.execute(
                "SELECT * FROM sales WHERE sale_date=? ORDER BY created_at", (day,)).fetchall()
            fname = f"sales_{day}.xlsx"
        else:
            sales_rows = conn.execute(
                "SELECT * FROM sales WHERE sale_date LIKE ? ORDER BY sale_date, created_at",
                (year + '-%',)).fetchall()
            fname = f"sales_{year}.xlsx"

        all_sales = []
        for s in sales_rows:
            prods = conn.execute("SELECT * FROM sale_products WHERE sale_id=?", (s['id'],)).fetchall()
            sale_dict = dict(s)
            sale_dict['payments'] = parse_payments(sale_dict.get('payments'))
            all_sales.append({'sale': sale_dict, 'products': [dict(p) for p in prods]})

        # Assign one receipt/invoice number per sale/order.
        # A sale is one customer order, no matter how many products/items it contains.
        # Format: RCPT-YYYY-####, where #### starts at 0001 for this exported year.
        receipt_counters = {}
        for entry in all_sales:
            sale_year = str(entry['sale'].get('sale_date') or year)[:4]
            receipt_counters[sale_year] = receipt_counters.get(sale_year, 0) + 1
            entry['sale']['receipt_invoice_no'] = f"RCPT-{sale_year}-{receipt_counters[sale_year]:04d}"

    wb = Workbook()
    wb.remove(wb.active)

    HDR_FILL = PatternFill("solid", fgColor="1A1916")
    HDR_FONT = Font(color="F7F6F2", bold=True, size=11)
    SUB_FILL = PatternFill("solid", fgColor="E2DFD6")
    SUB_FONT = Font(bold=True, color="1A1916")
    DAY_FILL = PatternFill("solid", fgColor="2F2D28")
    DAY_FONT = Font(color="F7F6F2", bold=True, size=12)
    TOTAL_FILL = PatternFill("solid", fgColor="E1F5EE")
    TOTAL_FONT = Font(bold=True, color="0F6E56")
    thin = Side(style='thin', color='D0CEC6')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    cols = ['Receipt/Invoice #','Product #','Customer','Product Type','Name on Product','Formula','Price ($)','Payment(s)','Sale Notes']

    def setup_sheet(ws):
        widths = [18, 10, 20, 16, 18, 40, 12, 28, 20]
        for idx, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        for ci, col in enumerate(cols, 1):
            cell = ws.cell(row=1, column=ci, value=col)
            cell.font = HDR_FONT
            cell.fill = HDR_FILL
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border
        ws.row_dimensions[1].height = 20
        ws.freeze_panes = 'A2'

    def formula_to_text(p):
        formula_str = ''
        if p.get('formula'):
            try:
                f = json.loads(p['formula']) if isinstance(p['formula'], str) else p['formula']
                formula_str = ' + '.join(f"{ing['pct']}% {ing['name']}" for ing in f)
            except Exception:
                formula_str = str(p['formula'])
        return formula_str

    def product_type_text(product):
        if product.get('product_type') == 'Custom':
            custom_name = (product.get('custom_product_name') or '').strip()
            return f'Custom: "{custom_name}"' if custom_name else 'Custom'
        return product.get('product_type') or ''

    def write_product_row(ws, row, item_no, sale, product, show_sale_info=False):
        # Keep the sheet compact: show sale-level payment/notes only once,
        # on the first product row of that sale. Do not add extra SALE SUMMARY rows.
        vals = [
            sale.get('receipt_invoice_no', '') if show_sale_info else '',
            item_no,
            sale['customer'],
            product_type_text(product),
            product.get('product_label') or '',
            formula_to_text(product),
            round(float(product.get('price') or 0), 2),
            payment_summary(sale.get('payments', [])) if show_sale_info else '',
            sale.get('notes','') if show_sale_info else ''
        ]
        for ci, v in enumerate(vals, 1):
            cell = ws.cell(row=row, column=ci, value=v)
            cell.border = border
            cell.alignment = Alignment(vertical='center', wrap_text=(ci in [6,8,9]))
            if ci == 7:
                cell.number_format = '#,##0.00'
            # Keep product numbers, customer names, and payment values normal-weight for readability.
            if show_sale_info and ci == 9 and v:
                cell.font = SUB_FONT
        ws.row_dimensions[row].height = 15

    def write_section_header(ws, row, label):
        # Strong visual separator so major sections are easy to distinguish.
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        cell = ws.cell(row=row, column=1, value=label)
        cell.font = DAY_FONT
        cell.fill = DAY_FILL
        cell.alignment = Alignment(horizontal='left', vertical='center')
        cell.border = border
        ws.row_dimensions[row].height = 22
        return row + 1

    def write_day_header(ws, row, sale_date):
        return write_section_header(ws, row, f"Sales Day: {sale_date}")

    def write_total_row(ws, row, label, amount):
        tr = ws.cell(row=row, column=6, value=label)
        tr.font = TOTAL_FONT
        tr.fill = TOTAL_FILL
        tr.border = border
        tr.alignment = Alignment(horizontal='right')
        tc = ws.cell(row=row, column=7, value=round(amount, 2))
        tc.font = TOTAL_FONT
        tc.fill = TOTAL_FILL
        tc.border = border
        tc.number_format = '#,##0.00'
        for ci in [1,2,3,4,5,8,9]:
            c = ws.cell(row=row, column=ci, value='')
            c.fill = TOTAL_FILL
            c.border = border
        return row + 1

    def product_display_name(product):
        if product.get('product_type') == 'Custom':
            custom_name = (product.get('custom_product_name') or '').strip()
            return f'Custom: "{custom_name}"' if custom_name else 'Custom'
        return product.get('product_type') or 'Unknown'

    def summarize_products(sales_data):
        summary = {}
        total_products = 0
        total_revenue = 0.0
        for entry in sales_data:
            for product in entry['products']:
                name = product_display_name(product)
                price = float(product.get('price') or 0)
                if name not in summary:
                    summary[name] = {'count': 0, 'revenue': 0.0}
                summary[name]['count'] += 1
                summary[name]['revenue'] += price
                total_products += 1
                total_revenue += price
        return summary, total_products, total_revenue

    def ingredient_summary_name(name):
        """Clean ingredient display name for summary tables.

        Example: Alth (Althair - Tier 3) -> Althair.
        The count still comes from each formula row, not from percentage amount.
        """
        text = (name or '').strip()
        if '(' in text and ')' in text:
            before = text.split('(', 1)[0].strip()
            inside = text.split('(', 1)[1].split(')', 1)[0].strip()
            main_inside = inside.split('-', 1)[0].strip()
            return main_inside or inside or before or text
        return text

    def summarize_ingredients(sales_data):
        """Count how many products used each ingredient.

        A product with 1% of an ingredient counts as 1 use.
        A product with 100% of an ingredient also counts as 1 use.
        Ingredients from the master INGREDIENTS list are included even if sold 0 times.
        """
        summary = {}
        for ingredient in INGREDIENTS:
            name = ingredient_summary_name(ingredient)
            if name and name.lower() != 'custom':
                summary.setdefault(name, 0)

        total_uses = 0
        for entry in sales_data:
            for product in entry['products']:
                formula = product.get('formula') or []
                if isinstance(formula, str):
                    try:
                        formula = json.loads(formula)
                    except Exception:
                        formula = []
                if not isinstance(formula, list):
                    continue

                # Count an ingredient only once per product even if it somehow appears twice.
                seen_in_product = set()
                for item in formula:
                    if not isinstance(item, dict):
                        continue
                    name = ingredient_summary_name(item.get('name'))
                    if not name or name in seen_in_product:
                        continue
                    seen_in_product.add(name)
                    summary[name] = summary.get(name, 0) + 1
                    total_uses += 1
        return summary, total_uses

    def payment_display_name(pay):
        if pay.get('method') == 'Custom' and pay.get('custom_method'):
            return str(pay.get('custom_method')).strip()
        return str(pay.get('method') or 'Unknown').strip() or 'Unknown'

    def summarize_payments(sales_data):
        summary = {}
        total_paid = 0.0
        for entry in sales_data:
            sale = entry['sale']
            sale_total = sum(float(p.get('price') or 0) for p in entry['products'])
            payments = parse_payments(sale.get('payments', []))
            if not payments:
                if sale_total:
                    summary.setdefault('Not recorded', 0.0)
                    summary['Not recorded'] += sale_total
                    total_paid += sale_total
                continue
            for pay in payments:
                name = payment_display_name(pay)
                try:
                    amount = float(pay.get('amount') or 0)
                except (TypeError, ValueError):
                    amount = 0.0
                # Old payment records may have a method but no amount. Put the full sale total
                # under that method so the payment table still reconciles to revenue.
                if amount == 0 and len(payments) == 1 and sale_total:
                    amount = sale_total
                summary.setdefault(name, 0.0)
                summary[name] += amount
                total_paid += amount
        return summary, total_paid

    def write_product_summary(ws, row, sales_data, title):
        product_summary, total_products, total_revenue = summarize_products(sales_data)
        payment_summary_data, total_paid = summarize_payments(sales_data)
        row += 1

        # Left table: product type totals. Right table: payment method revenue.
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        title_cell = ws.cell(row=row, column=1, value=title)
        title_cell.font = SUB_FONT
        title_cell.fill = SUB_FILL
        title_cell.alignment = Alignment(horizontal='left')
        title_cell.border = border

        ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=7)
        pay_title = ws.cell(row=row, column=5, value=title.replace('PRODUCT TOTALS', 'PAYMENT TOTALS'))
        pay_title.font = SUB_FONT
        pay_title.fill = SUB_FILL
        pay_title.alignment = Alignment(horizontal='left')
        pay_title.border = border
        row += 1

        product_headers = ['Product Type', 'Products Sold', 'Revenue ($)']
        payment_headers = ['Payment Method', 'Revenue ($)', '']
        for offset, header in enumerate(product_headers, start=1):
            cell = ws.cell(row=row, column=offset, value=header)
            cell.font = HDR_FONT
            cell.fill = HDR_FILL
            cell.alignment = Alignment(horizontal='center')
            cell.border = border
        for offset, header in enumerate(payment_headers, start=5):
            cell = ws.cell(row=row, column=offset, value=header)
            cell.font = HDR_FONT
            cell.fill = HDR_FILL
            cell.alignment = Alignment(horizontal='center')
            cell.border = border
        row += 1

        product_rows = []
        if not product_summary:
            product_rows.append(['No products sold in this period.', '', ''])
        else:
            for name, data in sorted(product_summary.items(), key=lambda x: (-x[1]['count'], x[0])):
                product_rows.append([name, data['count'], round(data['revenue'], 2)])

        payment_rows = []
        if not payment_summary_data:
            payment_rows.append(['No payments recorded.', '', ''])
        else:
            for name, amount in sorted(payment_summary_data.items(), key=lambda x: (-x[1], x[0])):
                payment_rows.append([name, round(amount, 2), ''])

        body_len = max(len(product_rows), len(payment_rows))
        for idx in range(body_len):
            if idx < len(product_rows):
                vals = product_rows[idx]
                for ci, v in enumerate(vals, 1):
                    cell = ws.cell(row=row, column=ci, value=v)
                    cell.border = border
                    if ci == 3 and isinstance(v, (int, float)):
                        cell.number_format = '#,##0.00'
            else:
                for ci in range(1, 4):
                    ws.cell(row=row, column=ci, value='').border = border

            if idx < len(payment_rows):
                vals = payment_rows[idx]
                for ci, v in zip([5, 6, 7], vals):
                    cell = ws.cell(row=row, column=ci, value=v)
                    cell.border = border
                    if ci == 6 and isinstance(v, (int, float)):
                        cell.number_format = '#,##0.00'
            else:
                for ci in range(5, 8):
                    ws.cell(row=row, column=ci, value='').border = border
            row += 1

        total_vals = ['TOTAL PRODUCTS SOLD', total_products, round(total_revenue, 2)]
        for ci, v in enumerate(total_vals, 1):
            cell = ws.cell(row=row, column=ci, value=v)
            cell.font = TOTAL_FONT
            cell.fill = TOTAL_FILL
            cell.border = border
            if ci == 3:
                cell.number_format = '#,##0.00'

        pay_total_vals = ['TOTAL PAYMENTS', round(total_paid, 2), '']
        for ci, v in zip([5, 6, 7], pay_total_vals):
            cell = ws.cell(row=row, column=ci, value=v)
            cell.font = TOTAL_FONT
            cell.fill = TOTAL_FILL
            cell.border = border
            if ci == 6:
                cell.number_format = '#,##0.00'
        return row + 1

    def write_ingredient_summary(ws, row, sales_data, title):
        ingredient_summary, total_uses = summarize_ingredients(sales_data)
        row += 1

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        title_cell = ws.cell(row=row, column=1, value=title)
        title_cell.font = SUB_FONT
        title_cell.fill = SUB_FILL
        title_cell.alignment = Alignment(horizontal='left')
        title_cell.border = border
        row += 1

        for ci, header in enumerate(['Ingredient', 'Times Sold'], 1):
            cell = ws.cell(row=row, column=ci, value=header)
            cell.font = HDR_FONT
            cell.fill = HDR_FILL
            cell.alignment = Alignment(horizontal='center')
            cell.border = border
        row += 1

        if ingredient_summary:
            for name, count in sorted(ingredient_summary.items(), key=lambda x: (-x[1], x[0])):
                ws.cell(row=row, column=1, value=name).border = border
                ws.cell(row=row, column=2, value=count).border = border
                row += 1
        else:
            ws.cell(row=row, column=1, value='No ingredients sold in this period.').border = border
            ws.cell(row=row, column=2, value='').border = border
            row += 1

        total_label = ws.cell(row=row, column=1, value='TOTAL INGREDIENT USES')
        total_label.font = TOTAL_FONT
        total_label.fill = TOTAL_FILL
        total_label.border = border
        total_count = ws.cell(row=row, column=2, value=total_uses)
        total_count.font = TOTAL_FONT
        total_count.fill = TOTAL_FILL
        total_count.border = border
        return row + 1

    def make_detail_sheet(wb, title, sales_data, final_total_label, summary_only=False, include_ingredient_summary=False):
        ws = wb.create_sheet(title=title[:31])
        setup_sheet(ws)
        row = 2
        grand_total = 0.0
        item_no = 1

        if not summary_only:
            from collections import defaultdict
            by_day = defaultdict(list)
            for entry in sales_data:
                by_day[entry['sale']['sale_date']].append(entry)

            if not by_day:
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
                ws.cell(row=row, column=1, value='No sales for this period.')
                row += 1
            else:
                for sale_date in sorted(by_day.keys()):
                    row = write_day_header(ws, row, sale_date)
                    day_total = 0.0
                    item_no = 1  # Product number resets for each sales day.
                    day_sales = by_day[sale_date]
                    for entry in day_sales:
                        sale = entry['sale']
                        products = entry['products']
                        sale_start_row = row
                        payment_text = payment_summary(sale.get('payments', []))
                        for product_index, product in enumerate(products):
                            write_product_row(ws, row, item_no, sale, product, show_sale_info=(product_index == 0))
                            item_no += 1
                            day_total += float(product.get('price') or 0)
                            row += 1

                        # If one sale has multiple products, make the payment easier to read by
                        # using one tall payment cell beside all products in that sale instead of
                        # making it look attached to only the first product row. Avoid merged
                        # cells when there is only one product because it adds no value.
                        sale_end_row = row - 1
                        if len(products) > 1:
                            # One receipt/invoice number and one customer name belong to
                            # the whole sale/order, so keep both as tall cells beside all
                            # products in that sale. Product # stays unmerged because each
                            # product/item gets its own number.
                            ws.merge_cells(start_row=sale_start_row, start_column=1, end_row=sale_end_row, end_column=1)
                            receipt_cell = ws.cell(row=sale_start_row, column=1)
                            receipt_cell.value = sale.get('receipt_invoice_no', '')
                            receipt_cell.alignment = Alignment(horizontal='left', vertical='center')
                            receipt_cell.border = border

                            ws.merge_cells(start_row=sale_start_row, start_column=3, end_row=sale_end_row, end_column=3)
                            customer_cell = ws.cell(row=sale_start_row, column=3)
                            customer_cell.value = sale.get('customer', '')
                            customer_cell.alignment = Alignment(horizontal='left', vertical='center')
                            customer_cell.border = border

                            for merged_row in range(sale_start_row + 1, sale_end_row + 1):
                                ws.cell(row=merged_row, column=1).border = border
                                ws.cell(row=merged_row, column=3).border = border

                        if len(products) > 1 and payment_text:
                            ws.merge_cells(start_row=sale_start_row, start_column=8, end_row=sale_end_row, end_column=8)
                            pay_cell = ws.cell(row=sale_start_row, column=8)
                            pay_cell.value = payment_text
                            pay_cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
                            pay_cell.border = border
                            for merged_row in range(sale_start_row + 1, sale_end_row + 1):
                                ws.cell(row=merged_row, column=8).border = border
                    row = write_total_row(ws, row, f"DAY TOTAL ({sale_date})", day_total)
                    row = write_product_summary(ws, row, day_sales, f"PRODUCT TOTALS ({sale_date})")
                    grand_total += day_total
                    row += 2  # Two blank rows between completed sales-day sections.

        if title == 'YEAR TOTAL':
            summary_label = f"YEAR SUMMARY — {year}"
        elif mode == 'day':
            summary_label = f"DAY SUMMARY — {title}"
        else:
            summary_label = f"MONTH SUMMARY — {title.upper()} {year}"

        row = write_section_header(ws, row, summary_label)
        row = write_product_summary(ws, row, sales_data, f"PRODUCT TOTALS - {final_total_label}")
        if include_ingredient_summary:
            row = write_ingredient_summary(ws, row, sales_data, f"INGREDIENT TOTALS - {final_total_label}")
        ws.auto_filter.ref = f"A1:I{max(1, row-1)}"
        return ws

    if mode == 'day':
        make_detail_sheet(wb, day, all_sales, 'DAY TOTAL')
    else:
        import calendar
        monthly_sales = {m: [] for m in range(1, 13)}
        for entry in all_sales:
            try:
                month_num = int(entry['sale']['sale_date'][5:7])
                monthly_sales[month_num].append(entry)
            except Exception:
                pass

        # Always create exactly 12 monthly sheets for the selected year.
        for month_num in range(1, 13):
            month_name = calendar.month_name[month_num]
            make_detail_sheet(wb, month_name, monthly_sales[month_num], f"{month_name.upper()} TOTAL")

        # Extra summary sheet for the full selected year. Keep it summary-only: no daily sales rows.
        make_detail_sheet(wb, 'YEAR TOTAL', all_sales, f'{year} YEAR TOTAL', summary_only=True, include_ingredient_summary=True)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=fname,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.route('/api/export/overall-revenue')
def export_overall_revenue():
    guard = require_admin_api()
    if guard: return guard
    """Export overall revenue in a clean multi-sheet workbook.

    Sheets:
    - Yearly Revenue: one summary row per year
    - Product Totals by Year: product totals grouped by year
    - Payment Totals by Year: payment totals grouped by year
    - All-Time Totals: all-time product totals, payment totals, and ingredient usage
    """
    with get_db() as conn:
        year_rows = conn.execute("""
            SELECT substr(s.sale_date, 1, 4) AS year,
                   COUNT(DISTINCT s.id) AS sales_count,
                   COUNT(sp.id) AS products_sold,
                   ROUND(COALESCE(SUM(sp.price), 0), 2) AS revenue
            FROM sales s
            JOIN sale_products sp ON sp.sale_id = s.id
            GROUP BY year
            ORDER BY year
        """).fetchall()
        product_rows = conn.execute("""
            SELECT substr(s.sale_date, 1, 4) AS year,
                   CASE WHEN sp.product_type='Custom' AND COALESCE(sp.custom_product_name,'') != '' THEN 'Custom: "' || sp.custom_product_name || '"' ELSE sp.product_type END AS product_name,
                   COUNT(sp.id) AS products_sold,
                   ROUND(COALESCE(SUM(sp.price), 0), 2) AS revenue
            FROM sales s
            JOIN sale_products sp ON sp.sale_id = s.id
            GROUP BY year, product_name
            ORDER BY year, products_sold DESC, product_name
        """).fetchall()
        all_time_product_rows = conn.execute("""
            SELECT CASE WHEN sp.product_type='Custom' AND COALESCE(sp.custom_product_name,'') != '' THEN 'Custom: "' || sp.custom_product_name || '"' ELSE sp.product_type END AS product_name,
                   COUNT(sp.id) AS products_sold,
                   ROUND(COALESCE(SUM(sp.price), 0), 2) AS revenue
            FROM sale_products sp
            GROUP BY product_name
            ORDER BY products_sold DESC, product_name
        """).fetchall()
        payment_rows = conn.execute("""
            SELECT substr(s.sale_date, 1, 4) AS year,
                   s.payments,
                   ROUND(COALESCE((SELECT SUM(sp.price) FROM sale_products sp WHERE sp.sale_id = s.id), 0), 2) AS sale_total
            FROM sales s
            ORDER BY s.sale_date
        """).fetchall()
        ingredient_rows = conn.execute("""
            SELECT sp.formula
            FROM sale_products sp
            WHERE sp.formula IS NOT NULL AND sp.formula != '' AND sp.formula != '[]'
        """).fetchall()

    wb = Workbook()
    ws_yearly = wb.active
    ws_yearly.title = 'Yearly Revenue'
    ws_products = wb.create_sheet(title='Product Totals by Year')
    ws_payments = wb.create_sheet(title='Payment Totals by Year')
    ws_all_time = wb.create_sheet(title='All-Time Totals')

    HDR_FILL = PatternFill("solid", fgColor="1A1916")
    HDR_FONT = Font(color="F7F6F2", bold=True, size=11)
    SUB_FILL = PatternFill("solid", fgColor="E2DFD6")
    SUB_FONT = Font(bold=True, color="1A1916")
    TOTAL_FILL = PatternFill("solid", fgColor="E1F5EE")
    TOTAL_FONT = Font(bold=True, color="0F6E56")
    thin = Side(style='thin', color='D0CEC6')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def payment_display_name_for_export(pay):
        if pay.get('method') == 'Custom' and pay.get('custom_method'):
            return str(pay.get('custom_method')).strip()
        return str(pay.get('method') or 'Unknown').strip() or 'Unknown'

    def write_headers(ws_obj, row, start_col, headers):
        for offset, header in enumerate(headers):
            cell = ws_obj.cell(row=row, column=start_col + offset, value=header)
            cell.font = HDR_FONT
            cell.fill = HDR_FILL
            cell.alignment = Alignment(horizontal='center')
            cell.border = border

    def style_total_row(ws_obj, row, start_col, values, money_offsets=()):
        for offset, v in enumerate(values):
            cell = ws_obj.cell(row=row, column=start_col + offset, value=v)
            cell.font = TOTAL_FONT
            cell.fill = TOTAL_FILL
            cell.border = border
            if offset in money_offsets:
                cell.number_format = '#,##0.00'

    def ingredient_summary_name_for_export(name):
        text_value = (name or '').strip()
        if '(' in text_value and ')' in text_value:
            before = text_value.split('(', 1)[0].strip()
            inside = text_value.split('(', 1)[1].split(')', 1)[0].strip()
            main_inside = inside.split('-', 1)[0].strip()
            return main_inside or inside or before or text_value
        return text_value

    def summarize_all_time_ingredients(rows):
        summary = {}
        ingredient_order = []

        # Start with every ingredient from the master list, including 0-sold ingredients.
        for ingredient in INGREDIENTS:
            name = ingredient_summary_name_for_export(ingredient)
            if name and name.lower() != 'custom' and name not in summary:
                summary[name] = 0
                ingredient_order.append(name)

        # Count each ingredient once per product, no matter whether it was 1% or 100%.
        for r in rows:
            try:
                formula = json.loads(r['formula']) if isinstance(r['formula'], str) else r['formula']
            except Exception:
                formula = []
            if not isinstance(formula, list):
                continue

            seen_in_product = set()
            for item in formula:
                if not isinstance(item, dict):
                    continue
                name = ingredient_summary_name_for_export(item.get('name'))
                if not name or name in seen_in_product:
                    continue
                seen_in_product.add(name)

                # Keep old/extra ingredients too, even if they are no longer in the master list.
                if name not in summary:
                    summary[name] = 0
                    ingredient_order.append(name)
                summary[name] += 1

        total_uses = sum(summary.values())
        return [(name, summary.get(name, 0)) for name in ingredient_order], total_uses

    all_time_ingredient_rows, all_time_ingredient_uses = summarize_all_time_ingredients(ingredient_rows)

    payment_by_year = {}
    all_time_payments = {}
    for r in payment_rows:
        year_key = r['year'] or 'Unknown'
        sale_total = float(r['sale_total'] or 0)
        payments = parse_payments(r['payments'])

        if not payments:
            if sale_total:
                payment_by_year.setdefault(year_key, {}).setdefault('Not recorded', 0.0)
                payment_by_year[year_key]['Not recorded'] += sale_total
                all_time_payments.setdefault('Not recorded', 0.0)
                all_time_payments['Not recorded'] += sale_total
            continue

        for pay in payments:
            name = payment_display_name_for_export(pay)
            try:
                amount = float(pay.get('amount') or 0)
            except (TypeError, ValueError):
                amount = 0.0

            # Old payment records may have a method but no amount.
            # If there is only one payment method, count the full sale total for that method.
            if amount == 0 and len(payments) == 1 and sale_total:
                amount = sale_total

            payment_by_year.setdefault(year_key, {}).setdefault(name, 0.0)
            payment_by_year[year_key][name] += amount
            all_time_payments.setdefault(name, 0.0)
            all_time_payments[name] += amount

    # Sheet 1: Yearly Revenue
    for idx, width in enumerate([14, 12, 16, 14], 1):
        ws_yearly.column_dimensions[get_column_letter(idx)].width = width
    write_headers(ws_yearly, 1, 1, ['Year', 'Sales', 'Products Sold', 'Revenue ($)'])

    total_revenue = 0.0
    total_sales = 0
    total_products = 0
    row_num = 2

    if not year_rows:
        ws_yearly.cell(row=row_num, column=1, value='No sales yet.').border = border
        row_num += 1
    else:
        for r in year_rows:
            year_key = r['year'] or 'Unknown'
            sales_count = int(r['sales_count'] or 0)
            products_sold = int(r['products_sold'] or 0)
            revenue = float(r['revenue'] or 0)
            total_sales += sales_count
            total_products += products_sold
            total_revenue += revenue

            vals = [year_key, sales_count, products_sold, round(revenue, 2)]
            for ci, v in enumerate(vals, 1):
                cell = ws_yearly.cell(row=row_num, column=ci, value=v)
                cell.border = border
                cell.alignment = Alignment(vertical='center')
                if ci == 4:
                    cell.number_format = '#,##0.00'
            row_num += 1

    style_total_row(ws_yearly, row_num, 1, ['OVERALL TOTAL', total_sales, total_products, round(total_revenue, 2)], money_offsets=(3,))

    # Sheet 2: Product Totals by Year — grouped format.
    # Shows the year only once per group instead of repeating it on every row.
    for idx, width in enumerate([14, 28, 16, 14], 1):
        ws_products.column_dimensions[get_column_letter(idx)].width = width
    write_headers(ws_products, 1, 1, ['Year', 'Product Type', 'Products Sold', 'Revenue ($)'])

    product_rows_by_year = {}
    for r in product_rows:
        product_rows_by_year.setdefault(r['year'] or 'Unknown', []).append(r)

    product_total_count = 0
    product_total_revenue = 0.0
    row_num = 2

    if not product_rows_by_year:
        ws_products.cell(row=row_num, column=1, value='No products sold.').border = border
        row_num += 1
    else:
        for year_key in sorted(product_rows_by_year.keys()):
            year_count = 0
            year_revenue = 0.0
            first_row_for_year = True

            for r in product_rows_by_year.get(year_key, []):
                products_sold = int(r['products_sold'] or 0)
                revenue = float(r['revenue'] or 0)
                year_count += products_sold
                year_revenue += revenue
                product_total_count += products_sold
                product_total_revenue += revenue

                vals = [
                    year_key if first_row_for_year else '',
                    r['product_name'],
                    products_sold,
                    round(revenue, 2)
                ]
                first_row_for_year = False

                for ci, v in enumerate(vals, 1):
                    cell = ws_products.cell(row=row_num, column=ci, value=v)
                    cell.border = border
                    if ci == 4:
                        cell.number_format = '#,##0.00'
                row_num += 1

            style_total_row(ws_products, row_num, 1, [f"{year_key} TOTAL", '', year_count, round(year_revenue, 2)], money_offsets=(3,))
            row_num += 2

    style_total_row(ws_products, row_num, 1, ['ALL YEARS TOTAL', '', product_total_count, round(product_total_revenue, 2)], money_offsets=(3,))

    # Sheet 3: Payment Totals by Year — grouped format.
    # Shows the year only once per group instead of repeating it on every row.
    for idx, width in enumerate([14, 26, 14], 1):
        ws_payments.column_dimensions[get_column_letter(idx)].width = width
    write_headers(ws_payments, 1, 1, ['Year', 'Payment Method', 'Revenue ($)'])

    payment_total_revenue = 0.0
    row_num = 2
    years_with_payments = sorted(payment_by_year.keys())

    if not years_with_payments:
        ws_payments.cell(row=row_num, column=1, value='No payments recorded.').border = border
        row_num += 1
    else:
        for year_key in years_with_payments:
            year_total = 0.0
            first_row_for_year = True

            for method, amount in sorted(payment_by_year[year_key].items(), key=lambda x: (-x[1], x[0])):
                year_total += amount
                payment_total_revenue += amount

                vals = [
                    year_key if first_row_for_year else '',
                    method,
                    round(amount, 2)
                ]
                first_row_for_year = False

                for ci, v in enumerate(vals, 1):
                    cell = ws_payments.cell(row=row_num, column=ci, value=v)
                    cell.border = border
                    if ci == 3:
                        cell.number_format = '#,##0.00'
                row_num += 1

            style_total_row(ws_payments, row_num, 1, [f"{year_key} TOTAL", '', round(year_total, 2)], money_offsets=(2,))
            row_num += 2

    style_total_row(ws_payments, row_num, 1, ['ALL YEARS TOTAL', '', round(payment_total_revenue, 2)], money_offsets=(2,))

    # Sheet 4: All-Time Totals
    for idx, width in enumerate([28, 16, 14, 4, 26, 14, 4, 32, 16], 1):
        ws_all_time.column_dimensions[get_column_letter(idx)].width = width
    write_headers(ws_all_time, 1, 1, ['Product Type', 'Products Sold', 'Revenue ($)'])
    write_headers(ws_all_time, 1, 5, ['Payment Method', 'Revenue ($)'])
    write_headers(ws_all_time, 1, 8, ['Ingredient', 'Times Sold'])

    product_rnum = 2
    total_count = 0
    total_rev = 0.0
    if all_time_product_rows:
        for r in all_time_product_rows:
            count = int(r['products_sold'] or 0)
            rev = float(r['revenue'] or 0)
            total_count += count
            total_rev += rev
            for ci, v in enumerate([r['product_name'], count, round(rev, 2)], 1):
                cell = ws_all_time.cell(row=product_rnum, column=ci, value=v)
                cell.border = border
                if ci == 3:
                    cell.number_format = '#,##0.00'
            product_rnum += 1
    else:
        ws_all_time.cell(row=product_rnum, column=1, value='No products sold.').border = border
        product_rnum += 1
    style_total_row(ws_all_time, product_rnum, 1, ['TOTAL PRODUCTS SOLD', total_count, round(total_rev, 2)], money_offsets=(2,))

    payment_rnum = 2
    payment_total = 0.0
    if all_time_payments:
        for method, amount in sorted(all_time_payments.items(), key=lambda x: (-x[1], x[0])):
            payment_total += amount
            for ci, v in zip([5, 6], [method, round(amount, 2)]):
                cell = ws_all_time.cell(row=payment_rnum, column=ci, value=v)
                cell.border = border
                if ci == 6:
                    cell.number_format = '#,##0.00'
            payment_rnum += 1
    else:
        ws_all_time.cell(row=payment_rnum, column=5, value='No payments recorded.').border = border
        payment_rnum += 1
    style_total_row(ws_all_time, payment_rnum, 5, ['TOTAL PAYMENTS', round(payment_total, 2)], money_offsets=(1,))

    ingredient_rnum = 2
    for name, count in all_time_ingredient_rows:
        ws_all_time.cell(row=ingredient_rnum, column=8, value=name).border = border
        ws_all_time.cell(row=ingredient_rnum, column=9, value=count).border = border
        ingredient_rnum += 1
    style_total_row(ws_all_time, ingredient_rnum, 8, ['TOTAL INGREDIENT USES', all_time_ingredient_uses])

    # Freeze/filter each sheet.
    for ws_obj in [ws_yearly, ws_products, ws_payments, ws_all_time]:
        ws_obj.freeze_panes = 'A2'

    ws_yearly.auto_filter.ref = f"A1:D{max(1, ws_yearly.max_row)}"
    ws_products.auto_filter.ref = f"A1:D{max(1, ws_products.max_row)}"
    ws_payments.auto_filter.ref = f"A1:C{max(1, ws_payments.max_row)}"
    ws_all_time.auto_filter.ref = f"A1:I{max(1, ws_all_time.max_row)}"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='overall_revenue_by_year.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


if __name__ == '__main__':
    print("\n✓ Sales Tracker running at http://localhost:5000")
    print("  Close this window to stop.\n")
    app.run(debug=False, port=5000)
